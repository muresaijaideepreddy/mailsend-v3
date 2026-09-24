"""Regression checks for source fidelity, using synthetic documents only."""
import io
from docx import Document
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from mail.document_import import validate_extraction, read_document, ExtractionMismatch

class SourceFidelityTests(SimpleTestCase):
    def fixture(self):
        lines = ['To: john@example.test', 'CC: copy@example.test', 'Subject: Source subject',
                 'Date: October 1, 2026', 'Hi John Smith,', 'First mandatory paragraph.',
                 'Ask admin@example.test for context.', 'Last mandatory paragraph.']
        row = dict(to='john@example.test', cc='copy@example.test', bcc='', subject='Source subject',
                   send_date='2026-10-01', start_line=1, end_line=8, body_start_line=5, body_end_line=8)
        return lines, row

    def test_wrong_source_fields_are_rejected(self):
        for changes in [dict(to='admin@example.test'), dict(to='copy@example.test', cc=''),
                        dict(send_date='2028-01-15'), dict(subject='First mandatory paragraph.'), dict(cc='')]:
            with self.subTest(changes=changes):
                lines, row = self.fixture()
                with self.assertRaises(ExtractionMismatch):
                    validate_extraction({'messages': [dict(row, **changes)]}, lines)

    def test_body_cannot_be_shortened_or_omitted(self):
        for changes in [dict(body_end_line=6), dict(body_start_line=0, body_end_line=0),
                        dict(end_line=6, body_end_line=6)]:
            with self.subTest(changes=changes):
                lines, row = self.fixture()
                with self.assertRaises(ExtractionMismatch):
                    validate_extraction({'messages': [dict(row, **changes)]}, lines)

    def test_zero_results_cannot_hide_a_source_message(self):
        lines, _ = self.fixture()
        with self.assertRaises(ExtractionMismatch):
            validate_extraction({'messages': []}, lines)

    def test_second_message_cannot_disappear(self):
        lines, row = self.fixture()
        with self.assertRaises(ExtractionMismatch):
            validate_extraction({'messages': [row]}, lines + [''] + lines)

    def test_complete_source_is_accepted(self):
        lines, row = self.fixture()
        actual = validate_extraction({'messages': [row]}, lines)[0]
        self.assertEqual(actual['body'], '\n'.join(lines[4:]))
        self.assertEqual(actual['cc'], 'copy@example.test')

    def test_display_name_does_not_change_mailbox_identity(self):
        lines, row = self.fixture()
        lines[0] = 'To: John Smith <john@example.test>'
        self.assertEqual(validate_extraction({'messages': [row]}, lines)[0]['to'], 'john@example.test')

    def test_explicit_boundaries_separate_background_notes(self):
        lines, row = self.fixture()
        source = ['Background notes', 'EMAIL DRAFT'] + lines + ['END EMAIL DRAFT', 'Other background notes']
        row.update(start_line=2, end_line=11, body_start_line=7, body_end_line=10)
        self.assertEqual(validate_extraction({'messages': [row]}, source)[0]['body'], '\n'.join(lines[4:]))

    def test_marked_draft_cannot_hide_a_shortened_body(self):
        lines, row = self.fixture()
        source = ['EMAIL DRAFT'] + lines + ['END EMAIL DRAFT', 'Background notes']
        row.update(start_line=1, end_line=7, body_start_line=6, body_end_line=7)
        with self.assertRaises(ExtractionMismatch):
            validate_extraction({'messages': [row]}, source)

    def test_missing_source_fields_remain_empty(self):
        lines = ['To:', 'Subject:', 'Date:', 'Hi John,', 'Original message.']
        row = dict(to='', cc='', bcc='', subject='', send_date='', start_line=1, end_line=5,
                   body_start_line=4, body_end_line=5)
        result = validate_extraction({'messages': [row]}, lines)[0]
        self.assertEqual(result['to'], '')
        self.assertIsNone(result['send_date'])

    def test_reader_includes_header_footer_and_nested_tables(self):
        for kind in ('header', 'footer', 'nested'):
            with self.subTest(kind=kind):
                doc = Document()
                doc.add_paragraph('Main text')
                text = 'Important source text'
                if kind == 'nested':
                    doc.add_table(rows=1, cols=1).cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = text
                else:
                    getattr(doc.sections[0], kind).paragraphs[0].text = text
                buf = io.BytesIO(); doc.save(buf)
                self.assertIn(text, read_document(SimpleUploadedFile('fixture.docx', buf.getvalue())))

    def test_merged_cells_are_not_duplicated(self):
        doc = Document(); table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = 'Unique body text'
        buf = io.BytesIO(); doc.save(buf)
        self.assertEqual(read_document(SimpleUploadedFile('merged.docx', buf.getvalue())).count('Unique body text'), 1)
