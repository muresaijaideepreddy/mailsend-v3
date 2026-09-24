"""Document import and Claude contract checks; no real API or mail delivery."""
import io
import json
from unittest.mock import MagicMock, patch

import requests
from docx import Document
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from pypdf import PdfWriter

from mail.document_import import extract_drafts, read_document, redact_credentials, validate_extraction
from mail.models import AuditEvent, DocumentImport, Message
from mail.services import send_message
from mail.tests.test_security import WorkspaceTestCase

TEXT = 'To: jane@example.test\nSubject: Hello\nDate: 2026-09-23\nHi Jane,\nPlease review the report.'
ROW = dict(to='jane@example.test', cc='', bcc='', subject='Hello', body='Hi Jane,\nPlease review the report.', send_date='2026-09-23', start_line=1, end_line=5)


def upload(text=TEXT):
    doc = Document()
    doc.add_paragraph(text)
    stream = io.BytesIO()
    doc.save(stream)
    return SimpleUploadedFile('messages.docx', stream.getvalue())


def response_for(row=None, **overrides):
    payload = {'stop_reason': 'end_turn', 'content': [{'type': 'text', 'text': json.dumps({'messages': [ROW if row is None else row]})}]}
    payload.update(overrides)
    response = MagicMock()
    response.__enter__.return_value = response
    response.status_code = 200
    response.iter_content.return_value = [json.dumps(payload).encode()]
    return response


class DocumentParserTests(SimpleTestCase):
    def test_import_recipient_separators_and_missing_labels(self):
        from mail.document_import import parse_import_addresses
        for value in ('jane@example.test; bob@example.test', 'jane@example.test\nbob@example.test'):
            self.assertEqual(parse_import_addresses(value), 'jane@example.test, bob@example.test')
        for value in ('N/A', 'Not provided', 'TBD', ''):
            self.assertEqual(parse_import_addresses(value), '')
        with self.assertRaises(ValidationError):
            parse_import_addresses('jane@example.test; invalid-address')

    def test_invalid_import_recipient_keeps_draft_and_requires_review(self):
        rows = validate_extraction({'messages': [dict(ROW, cc='not-an-address')]}, TEXT.splitlines())
        self.assertEqual(rows[0]['to'], '')
        self.assertEqual(rows[0]['cc'], '')
        self.assertEqual(rows[0]['body'], ROW['body'])
        self.assertIn('CC: not-an-address', rows[0]['imported_recipient_notes'])
        self.assertIn('TO: jane@example.test', rows[0]['imported_recipient_notes'])

    def test_body_range_copies_every_source_line(self):
        lines = TEXT.splitlines() + ['A detail that must not disappear.', '', 'Original sign-off.']
        row = {key: value for key, value in ROW.items() if key != 'body'}
        row.update(end_line=len(lines), body_start_line=4, body_end_line=len(lines))
        result = validate_extraction({'messages': [row]}, lines)
        self.assertEqual(result[0]['body'], '\n'.join(lines[3:]))

    def test_body_range_rejects_outside_message_and_credentials(self):
        row = {key: value for key, value in ROW.items() if key != 'body'}
        for first, last in ((0, 4), (4, 6), (True, 5), (5, 4)):
            with self.subTest(first=first, last=last), self.assertRaises(ValueError):
                validate_extraction({'messages': [dict(row, body_start_line=first, body_end_line=last)]}, TEXT.splitlines())
        with self.assertRaises(ValueError):
            validate_extraction({'messages': [dict(row, body_start_line=4, body_end_line=5)]},
                                TEXT.splitlines()[:3] + ['Hi Jane,', '[credential removed]'])

    def test_absent_body_range_stays_blank(self):
        row = {key: value for key, value in ROW.items() if key != 'body'}
        row.update(body_start_line=0, body_end_line=0, end_line=3)
        self.assertEqual(validate_extraction({'messages': [row]}, TEXT.splitlines()[:3])[0]['body'], '')

    @override_settings(MAILSEND_CLAUDE_API_KEY='synthetic', MAILSEND_CLAUDE_PROVIDER='anthropic')
    def test_mismatch_reports_rule_without_leaking_content(self):
        changed = dict(ROW, body='Private invented content')
        with patch('mail.document_import.requests.post', return_value=response_for(changed)):
            with self.assertRaises(ValidationError) as caught:
                extract_drafts(TEXT)
        self.assertIn('the body was changed', str(caught.exception))
        self.assertNotIn('Private invented content', str(caught.exception))

    def test_word_breaks_and_table_order(self):
        doc = Document()
        doc.add_paragraph('To:\nSubject: First')
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'Table message'
        doc.add_paragraph('Last')
        stream = io.BytesIO(); doc.save(stream)
        text = read_document(SimpleUploadedFile('test.docx', stream.getvalue()))
        self.assertEqual(text, 'To:\nSubject: First\nTable message\nLast')

    def test_text_pdf_preserves_text(self):
        from pypdf.generic import DecodedStreamObject, NameObject, DictionaryObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 50 700 Td (Hello from PDF) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        output = io.BytesIO(); writer.write(output)
        self.assertIn('Hello from PDF', read_document(SimpleUploadedFile('text.pdf', output.getvalue())))

    def test_secret_values_are_removed_before_processing(self):
        text = redact_credentials('Keep\nPassword: secret-one\nAPI key:\nsecret-two\nDone')
        self.assertNotIn('secret-one', text)
        self.assertNotIn('secret-two', text)
        self.assertIn('Keep', text)
        self.assertIn('Done', text)

    def test_corrupt_and_unsupported_uploads(self):
        for name in ('broken.docx', 'broken.pdf', 'old.doc', 'script.exe'):
            with self.subTest(name=name), self.assertRaises(ValidationError):
                read_document(SimpleUploadedFile(name, b'not a document'))

    def test_large_upload_rejected(self):
        with self.assertRaises(ValidationError):
            read_document(SimpleUploadedFile('large.docx', b'x' * (5 * 1024 * 1024 + 1)))

    def test_empty_and_encrypted_pdf(self):
        for encrypted in (False, True):
            writer = PdfWriter(); writer.add_blank_page(width=72, height=72)
            if encrypted:
                writer.encrypt('synthetic-password')
            stream = io.BytesIO(); writer.write(stream)
            with self.subTest(encrypted=encrypted), self.assertRaises(ValidationError):
                read_document(SimpleUploadedFile('test.pdf', stream.getvalue()))

    def test_pdf_page_limit(self):
        writer = PdfWriter()
        for _ in range(101):
            writer.add_blank_page(width=72, height=72)
        stream = io.BytesIO(); writer.write(stream)
        with self.assertRaisesMessage(ValidationError, '100 pages'):
            read_document(SimpleUploadedFile('test.pdf', stream.getvalue()))


@override_settings(MAILSEND_CLAUDE_API_KEY='synthetic-test-key', MAILSEND_CLAUDE_MODEL='claude-sonnet-4-6')
class ClaudeContractTests(SimpleTestCase):
    @patch('mail.document_import.requests.post')
    def test_structured_request_and_valid_result(self, post):
        post.return_value = response_for()
        rows = extract_drafts(TEXT)
        self.assertEqual(rows[0]['to'], ROW['to'])
        self.assertEqual(rows[0]['body'], ROW['body'])
        args, kwargs = post.call_args
        self.assertEqual(args[0], 'https://api.anthropic.com/v1/messages')
        self.assertFalse(kwargs['allow_redirects'])
        self.assertEqual(kwargs['json']['output_config']['format']['type'], 'json_schema')

    @override_settings(MAILSEND_CLAUDE_API_KEY='')
    @patch('mail.document_import.requests.post')
    def test_blank_key_never_calls_api(self, post):
        with self.assertRaisesMessage(ValidationError, 'not configured'):
            extract_drafts(TEXT)
        post.assert_not_called()

    @patch('mail.document_import.requests.post')
    def test_timeout_and_rejected_request(self, post):
        post.side_effect = requests.Timeout('sensitive provider details')
        with self.assertRaisesMessage(ValidationError, 'could not be reached'):
            extract_drafts(TEXT)
        post.side_effect = None
        post.return_value = response_for(); post.return_value.status_code = 401
        with self.assertRaisesMessage(ValidationError, 'Check the API key'):
            extract_drafts(TEXT)

    @patch('mail.document_import.requests.post')
    def test_refusal_truncation_and_malformed_json(self, post):
        for data in ({'stop_reason': 'max_tokens'}, {'stop_reason': 'refusal'}, {'content': [{'type': 'text', 'text': 'invalid'}]}):
            post.return_value = response_for(**data)
            with self.subTest(data=data), self.assertRaises(ValidationError):
                extract_drafts(TEXT)

    def test_invented_recipient_and_body_and_duplicate_spans_rejected(self):
        for changes in ({'to': 'invented@example.test'}, {'body': 'Invented body'}, {'start_line': 0}, {'end_line': 99}):
            with self.subTest(changes=changes), self.assertRaises((ValueError, ValidationError)):
                validate_extraction({'messages': [dict(ROW, **changes)]}, TEXT.splitlines())
        with self.assertRaises(ValueError):
            validate_extraction({'messages': [ROW, ROW]}, TEXT.splitlines())


@override_settings(MAILSEND_CLAUDE_API_KEY='synthetic-test-key')
class DocumentWorkflowTests(WorkspaceTestCase):
    def token(self):
        return self.client.get(reverse('mail:document_upload')).context['form'].initial['token']

    @patch('mail.document_import.requests.post')
    def test_upload_directly_creates_shared_draft_and_redacts(self, post):
        post.return_value = response_for()
        self.login()
        result = self.client.post(reverse('mail:document_upload'), {'document': upload(TEXT + '\nPassword: do-not-transmit'), 'token': self.token()})
        self.assertEqual(result.status_code, 302)
        draft = Message.objects.get(imported_from='messages.docx')
        self.assertEqual(draft.created_by, self.executive)
        self.assertEqual(draft.status, 'draft')
        self.assertFalse(draft.attachments.exists())
        self.assertNotIn('do-not-transmit', json.dumps(post.call_args.kwargs['json']))
        self.assertEqual(DocumentImport.objects.get().status, 'complete')
        self.login(self.assistant)
        self.assertContains(self.client.get(reverse('mail:dashboard')), 'Hello')
        self.assertEqual(self.client.get(reverse('mail:edit', args=[draft.pk])).status_code, 200)
        self.assert_denied(self.client.post(reverse('mail:delete', args=[draft.pk]), {'version': 1}))
        self.assert_denied(self.client.get(reverse('mail:send', args=[draft.pk])))
        self.login(self.outsider)
        self.assert_denied(self.client.get(reverse('mail:detail', args=[draft.pk])))

    @patch('mail.document_import.requests.post')
    def test_invalid_cc_import_saved_with_review_note_and_send_blocked(self, post):
        post.return_value = response_for(dict(ROW, cc='invalid-address'))
        self.login()
        result = self.client.post(reverse('mail:document_upload'), {'document': upload(), 'token': self.token()})
        self.assertEqual(result.status_code, 302)
        draft = Message.objects.get(imported_from='messages.docx')
        self.assertIn('CC: invalid-address', draft.imported_recipient_notes)
        self.assertEqual(draft.body, ROW['body'])
        self.assertFalse(draft.ready_to_send)
        with patch('mail.services.build_email') as build, self.assertRaises(ValidationError):
            send_message(self.executive, draft.pk, expected_version=1)
        build.assert_not_called()
        self.login(self.assistant)
        self.assertContains(self.client.get(reverse('mail:edit', args=[draft.pk])), 'CC: invalid-address')

    @patch('mail.document_import.requests.post')
    def test_missing_fields_saved_editable_but_never_sent(self, post):
        post.return_value = response_for(dict(ROW, to='', subject='', send_date=''))
        self.login()
        self.client.post(reverse('mail:document_upload'), {'document': upload('To:\nSubject:\nDate:\nHi Jane,\nPlease review the report.'), 'token': self.token()})
        draft = Message.objects.get(imported_from='messages.docx')
        self.assertEqual(draft.to, '')
        self.assertIsNone(draft.send_date)
        self.assertFalse(draft.ready_to_send)
        with patch('mail.services.build_email') as build, self.assertRaises(ValidationError):
            send_message(self.executive, draft.pk, expected_version=1)
        build.assert_not_called()
        self.login(self.assistant)
        result = self.client.post(reverse('mail:edit', args=[draft.pk]), self.draft_data(draft, send_date='', subject='', body='Worker revision'))
        self.assertEqual(result.status_code, 302)
        draft.refresh_from_db(); self.assertEqual(draft.version, 2)
        self.assertEqual(draft.body, 'Worker revision')
        self.assertTrue(AuditEvent.objects.filter(message=draft, actor=self.assistant, action='draft.updated').exists())
        result = self.client.post(reverse('mail:edit', args=[draft.pk]), self.draft_data(draft, to='real@example.test', subject='Completed', send_date=timezone.localdate()))
        self.assertEqual(result.status_code, 302)
        draft.refresh_from_db(); self.assertTrue(draft.ready_to_send)

    @patch('mail.document_import.requests.post')
    def test_duplicate_submit_calls_api_once(self, post):
        post.return_value = response_for(); self.login(); token = self.token()
        file_bytes = upload().read()
        for _ in range(2):
            self.client.post(reverse('mail:document_upload'), {'document': SimpleUploadedFile('messages.docx', file_bytes), 'token': token})
        self.assertEqual(post.call_count, 1)
        self.assertEqual(Message.objects.filter(imported_from='messages.docx').count(), 1)

    @patch('mail.document_import.requests.post')
    def test_failed_upload_can_retry_with_different_file_without_reload(self, post):
        self.login(); first_token = self.token()
        post.side_effect = requests.Timeout('synthetic timeout')
        failed = self.client.post(reverse('mail:document_upload'), {'document': upload(), 'token': first_token})
        self.assertEqual(failed.status_code, 200)
        next_token = failed.context['form']['token'].value()
        self.assertNotEqual(next_token, first_token)
        post.side_effect = None; post.return_value = response_for()
        different = upload(); different.name = 'different-document.docx'
        completed = self.client.post(reverse('mail:document_upload'), {'document': different, 'token': next_token})
        self.assertEqual(completed.status_code, 302)
        self.assertEqual(DocumentImport.objects.filter(status='failed').count(), 1)
        self.assertEqual(DocumentImport.objects.filter(status='complete').count(), 1)
        self.assertEqual(Message.objects.filter(imported_from='different-document.docx').count(), 1)

    @patch('mail.document_import.requests.post')
    def test_duplicate_error_renders_fresh_form_for_next_upload(self, post):
        self.login(); old_token = self.token(); post.return_value = response_for()
        file_bytes = upload().read()
        self.client.post(reverse('mail:document_upload'), {'document': SimpleUploadedFile('messages.docx', file_bytes), 'token': old_token})
        duplicate = self.client.post(reverse('mail:document_upload'), {'document': SimpleUploadedFile('messages.docx', file_bytes), 'token': old_token})
        self.assertEqual(post.call_count, 1)
        new_token = duplicate.context['form']['token'].value()
        self.assertNotEqual(new_token, old_token)
        fresh = upload(); fresh.name = 'next-upload.docx'
        result = self.client.post(reverse('mail:document_upload'), {'document': fresh, 'token': new_token})
        self.assertEqual(result.status_code, 302)
        self.assertEqual(post.call_count, 2)
        self.assertEqual(Message.objects.filter(imported_from='next-upload.docx').count(), 1)

    @patch('mail.document_import.requests.post')
    def test_different_files_work_sequentially_even_with_cached_form_token(self, post):
        self.login(); token = self.token(); post.return_value = response_for()
        for filename, text in [('first.docx', TEXT), ('second.docx', TEXT.replace('report.', 'second report.'))]:
            post.return_value = response_for(dict(ROW, body='\n'.join(text.splitlines()[3:])))
            document = upload(text); document.name = filename
            result = self.client.post(reverse('mail:document_upload'), {'document': document, 'token': token})
            self.assertEqual(result.status_code, 302)
        self.assertEqual(post.call_count, 2)
        self.assertEqual(DocumentImport.objects.filter(status='complete').count(), 2)

    @patch('mail.document_import.requests.post')
    def test_changed_content_with_same_filename_and_cached_token_is_new_upload(self, post):
        self.login(); token = self.token(); post.return_value = response_for()
        for text in (TEXT, TEXT.replace('report.', 'second report.')):
            post.return_value = response_for(dict(ROW, body='\n'.join(text.splitlines()[3:])))
            result = self.client.post(reverse('mail:document_upload'), {'document': upload(text), 'token': token})
            self.assertEqual(result.status_code, 302)
        self.assertEqual(post.call_count, 2)

    @patch('mail.document_import.requests.post')
    def test_worker_and_forged_tokens_cannot_import(self, post):
        self.login(); token = self.token()
        self.login(self.outsider)
        result = self.client.post(reverse('mail:document_upload'), {'document': upload(), 'token': token})
        self.assertEqual(result.status_code, 200)
        self.login(self.assistant)
        self.assert_denied(self.client.get(reverse('mail:document_upload')))
        self.assert_denied(self.client.post(reverse('mail:document_upload'), {'document': upload(), 'token': token}))
        post.assert_not_called()

    @override_settings(MAILSEND_CLAUDE_API_KEY='')
    @patch('mail.document_import.requests.post')
    def test_disabled_ui_and_post(self, post):
        self.login()
        self.assertContains(self.client.get(reverse('mail:document_upload')), 'waiting for the Claude API key')
        self.client.post(reverse('mail:document_upload'), {'document': upload(), 'token': self.token()})
        post.assert_not_called(); self.assertFalse(DocumentImport.objects.exists())

    @patch('mail.document_import.extract_drafts')
    def test_invalid_second_draft_rolls_back_first(self, extract):
        values = {key: value for key, value in ROW.items() if not key.endswith('_line')}
        extract.return_value = [values, dict(values, subject='x' * 256)]
        self.login()
        self.client.post(reverse('mail:document_upload'), {'document': upload(), 'token': self.token()})
        self.assertFalse(Message.objects.filter(imported_from='messages.docx').exists())
        self.assertEqual(DocumentImport.objects.get().status, 'failed')

    def test_batch_preflight_blocks_before_any_delivery(self):
        self.make_message(self.executive, self.workspace, 'Incomplete', imported_from='file.docx', to='')
        self.login()
        token = self.client.get(reverse('mail:dashboard')).context['dashboard_batch_token']
        with patch('mail.views.send_message') as send:
            self.client.post(reverse('mail:send_current'), {'dashboard_token': token})
        send.assert_not_called()

    def test_manual_compose_still_requires_fields(self):
        self.login()
        result = self.client.post(reverse('mail:compose'), self.draft_data(to='', subject='', body='', send_date=''))
        self.assertEqual(result.status_code, 200)
        for name in ('to', 'subject', 'body', 'send_date'):
            self.assertIn(name, result.context['form'].errors)


class ImportDeploymentConfigTests(SimpleTestCase):
    def test_private_pythonanywhere_config_accepts_blank_and_configured_claude(self):
        import os
        import tempfile
        from dotenv import set_key
        from mail.tests.test_pythonanywhere import PythonAnywhereDeploymentTests
        helper = PythonAnywhereDeploymentTests()
        helper.setUp()
        try:
            helper.configured_file()
            with patch.dict(os.environ):
                values = helper.load()
                self.assertEqual(values['MAILSEND_CLAUDE_API_KEY'], '')
            set_key(str(helper.filename), 'MAILSEND_CLAUDE_API_KEY', 'synthetic-key')
            with patch.dict(os.environ):
                self.assertEqual(helper.load()['MAILSEND_CLAUDE_API_KEY'], 'synthetic-key')
        finally:
            helper.doCleanups()


@override_settings(MAILSEND_CLAUDE_PROVIDER='tamu', MAILSEND_CLAUDE_API_KEY='synthetic-key', MAILSEND_CLAUDE_MODEL='protected.Claude Sonnet 4.6')
class TamuContractTests(SimpleTestCase):
    @patch('mail.document_import.requests.post')
    def test_tamu_auth_payload_and_extraction(self, post):
        response = response_for()
        response.iter_content.return_value = [json.dumps({'choices': [{'finish_reason':'stop', 'message':{'content':json.dumps({'messages':[ROW]})}}]}).encode()]
        post.return_value = response
        self.assertEqual(extract_drafts(TEXT)[0]['to'], ROW['to'])
        args, kwargs = post.call_args
        self.assertEqual(args[0], 'https://chat-api.tamu.ai/openai/chat/completions')
        self.assertEqual(kwargs['headers']['Authorization'], 'Bearer synthetic-key')
        self.assertEqual(kwargs['json']['thinking'], {'type':'disabled'})
        self.assertNotIn('x-api-key', kwargs['headers'])

    @patch('mail.document_import.requests.post')
    def test_http_200_gateway_error_is_not_success(self, post):
        response = response_for()
        response.iter_content.return_value = [b'{"error":{"message":"synthetic provider error"}}']
        post.return_value = response
        with self.assertRaisesMessage(ValidationError, 'gateway rejected'):
            extract_drafts(TEXT)

    @patch('mail.document_import.requests.post')
    def test_tamu_truncation_rejected(self, post):
        response = response_for()
        response.iter_content.return_value = [json.dumps({'choices':[{'finish_reason':'length','message':{'content':json.dumps({'messages':[ROW]})}}]}).encode()]
        post.return_value = response
        with self.assertRaisesMessage(ValidationError, 'did not finish'):
            extract_drafts(TEXT)


class ImportTextCleanupTests(SimpleTestCase):
    def test_normal_password_discussion_is_preserved(self):
        text = 'Subject: Password reset reminder\nPlease reset your password using account settings.\nNo password or access token is included in this message.'
        self.assertEqual(redact_credentials(text), text)

    def test_actual_credential_labels_and_values_are_removed(self):
        text = 'Password: synthetic-secret\nAPI key\nsynthetic-key-value\nClient secret = another-secret\nNormal email body'
        cleaned = redact_credentials(text)
        for secret in ('synthetic-secret', 'synthetic-key-value', 'another-secret'):
            self.assertNotIn(secret, cleaned)
        self.assertIn('Normal email body', cleaned)

    def test_cross_page_furniture_is_removed_without_body_rewriting(self):
        from mail.document_import import clean_pdf_pages
        text = clean_pdf_pages(['To: a@example.test\nSubject: Update\nHi,\nFirst paragraph.\nTest notes | Page 1 of 2',
                                'Long message continues\nSecond paragraph.\nThank you.\nTest notes | Page 2 of 2'], [{'Test notes | Page 1 of 2'}, {'Long message continues', 'Test notes | Page 2 of 2'}])
        self.assertIn('First paragraph.\nSecond paragraph.', text)
        self.assertNotIn('Page 1', text)
        self.assertNotIn('Long message continues', text)
        row = dict(ROW, to='a@example.test', subject='Update', body='Hi,\nFirst paragraph.\nSecond paragraph.\nThank you.', send_date='', start_line=1, end_line=len(text.splitlines()))
        self.assertEqual(validate_extraction({'messages':[row]}, text.splitlines())[0]['body'], row['body'])

    def test_page_like_text_inside_body_is_preserved(self):
        from mail.document_import import clean_pdf_pages
        text = 'To: a@example.test\nSubject: Update\nHello,\nPage 1 of 2\nThe discussion continues\nAnother line\nThanks'
        self.assertEqual(clean_pdf_pages([text]), text)


class PageEdgePreservationTests(SimpleTestCase):
    def test_body_at_page_edge_is_not_removed_without_layout_evidence(self):
        from mail.document_import import clean_pdf_pages
        pages = ['Hi,\nPage 1 of 2', 'The discussion continues\nThank you.']
        self.assertEqual(clean_pdf_pages(pages), '\n'.join(pages))
